"""The verification report (PLAN-VERIFICATION-REPORT "Report structure"), as a Word document in the
style of the analysis report (Arial, shaded table headers, header and footer bands)."""
from __future__ import annotations

import datetime as dt
import hashlib
import json
import platform
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt

from ..report import BLUE, GREY, _header_footer, _kv_table, _para, _picture, _style_document, _table
from . import text as T

PASS_FAIL = {True: "PASS", False: "FAIL"}
USER_NOTE = T.USER_NOTE
NO_SWS_BUCKLING = ("Not available: SOLIDWORKS Premium has no buckling study (linear buckling needs the separately "
                   "licensed SOLIDWORKS Simulation Professional)")
SOURCES = [
    ("CalculiX home page", "Authors, history, licence, releases", "https://www.calculix.de"),
    ("CalculiX 2.22 manual, section 11, verification examples", "The test examples that ship with the solver",
     "https://www.dhondt.de/ccx_2.22.pdf"),
    ("Official examples repository", "Official test examples, including the structural set",
     "https://github.com/calculix/examples"),
    ("NAFEMS LE1, elliptic membrane (bConverged)", "CalculiX 1.5 against the NAFEMS target and ANSYS",
     "http://www.bconverged.com/benchmarks/le1.php"),
    ("NAFEMS LE10, thick plate (bConverged)", "The same comparison for a 3-D solid benchmark",
     "http://www.bconverged.com/benchmarks/le10.php"),
    ("Parametric NAFEMS LE10 study (seamplex)", "CalculiX among several solvers converging on LE10 (linked report)",
     "https://seamplex.com/feenox/tests/nafems/le10/"),
    ("Independent validation suite", "Third-party decks with closed-form references (small, informal)",
     "https://github.com/larsmei/calculix-validation"),
    ("Theory reference", "G. Dhondt, The Finite Element Method for Three-Dimensional Thermomechanical Applications, "
                         "Wiley, 2004, ISBN 0-470-85752-8", ""),
    ("NAFEMS on code verification", "Why benchmark testing is the accepted way to verify an analysis code",
     "https://www.nafems.org/community/the-analysis-agenda/code-verification/"),
    ("SOLIDWORKS Simulation product page", "Which analyses each SOLIDWORKS package includes",
     "https://www.solidworks.com/product/solidworks-simulation"),
]


def _pct(x) -> str:
    """Signed percentage; three decimals when the difference is below 0.1 %, so small is not shown as 0."""
    if x is None:
        return "-"
    if abs(x) < 5e-7:
        return "0.000 %"
    return f"{x * 100:+.3f} %" if abs(x) < 1e-3 else f"{x * 100:+.2f} %"


def _num(x) -> str:
    """Five significant figures; thousands separated."""
    if x is None:
        return "-"
    if abs(x) >= 1e5:
        return f"{x:,.0f}"
    return f"{x:,.5g}" if abs(x) >= 1000 else f"{x:.5g}"


def _sha256(p: Path) -> str:
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def _diagram(path: Path) -> Path:
    """How the analysis works: one row of boxes, CalculiX highlighted."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch
    boxes = [("SolidWorks\nassembly", "#e8eef7"), ("AsmFEA add-in\ngeometry, materials,\nconnections, loads", "#e8eef7"),
             ("AsmFEA pipeline\nmesh + CalculiX\ninput file", "#e8eef7"), ("CalculiX\nsolves the\nmechanics", "#ffe2b3"),
             ("AsmFEA\nreads results back:\nchecks, plots, report", "#e8eef7")]
    fig, ax = plt.subplots(figsize=(10, 2.2), dpi=180)
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 2.2)
    ax.axis("off")
    w, gap = 1.6, 0.4
    for i, (text, colour) in enumerate(boxes):
        x = 0.2 + i * (w + gap)
        ax.add_patch(FancyBboxPatch((x, 0.45), w, 1.3, boxstyle="round,pad=0.04,rounding_size=0.12",
                                    fc=colour, ec="#1f4e79", lw=1.4))
        ax.text(x + w / 2, 1.1, text, ha="center", va="center", fontsize=8.5, family="Arial")
        if i < len(boxes) - 1:
            ax.annotate("", xy=(x + w + gap - 0.04, 1.1), xytext=(x + w + 0.04, 1.1),
                        arrowprops=dict(arrowstyle="-|>", color="#1f4e79", lw=1.4))
    ax.text(5, 0.12, "The add-in prepares the model and reports the results; the stresses are calculated by CalculiX.",
            ha="center", fontsize=8, style="italic", family="Arial", color="#444444")
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def _mesh_plot(rows: list[dict], path: Path) -> Path:
    """C1: difference from theory of each read-out against degrees of freedom, with the 5 % band."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=(9, 3.4), dpi=170)
    ax.axhspan(-5, 5, color="#e8f4ea", label="acceptance band (5 %)")
    ax.axhline(0, color="#d9412b", ls="--", lw=1.2, label="theory")
    for q, colour in (("Tip deflection", "#1f6fd1"), ("Bending stress at mid-length", "#f28a1a")):
        pts = sorted((m["measured"]["dof"], next(r for r in m["rows"] if r["quantity"] == q)["difference"] * 100,
                      m["fidelity"]) for m in rows)
        ax.plot([p[0] for p in pts], [p[1] for p in pts], "o-", color=colour, label=q)
        for dof, d, fid in (pts if colour == "#1f6fd1" else []):
            ax.annotate(fid, (dof, d), textcoords="offset points", xytext=(4, 6), fontsize=7.5, color="#444444")
    ax.set_xscale("log")
    ax.set_ylim(-6, 6)
    ax.set_xlabel("degrees of freedom (reported mesh)")
    ax.set_ylabel("AsmFEA - theory (%)")
    ax.grid(alpha=0.3)
    ax.legend(fontsize=7.5, loc="lower right", ncol=2)
    fig.tight_layout()
    fig.savefig(path, facecolor="white")
    plt.close(fig)
    return path


def _theory_text(case_dir: Path) -> str:
    r = subprocess.run([sys.executable, "theory.py"], cwd=case_dir, capture_output=True, text=True,
                       encoding="utf-8", errors="replace")
    return r.stdout.strip()


def _mono(doc, text: str, size: float = 7.5) -> None:
    """Monospaced lines (a worked calculation), no space between them."""
    for line in text.splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        run = p.add_run(line or " ")
        run.font.name = "Consolas"
        run.font.size = Pt(size)


def build(result: dict, root: Path, out: Path, versions: dict) -> Path:
    """result: verify.evaluate(); versions: asmfea and solver versions for the stamp."""
    root, out = Path(root), Path(out)
    work = out.parent / (out.stem + "_files")
    work.mkdir(parents=True, exist_ok=True)
    doc = Document()
    _style_document(doc)
    when = dt.date.today()
    stamp = f"AsmFEA {versions['asmfea']}  |  CalculiX {versions['calculix']}"
    _header_footer(doc, "AsmFEA verification report", f"{when:%d %B %Y}", stamp)
    rows = result["rows"]
    sws_any = any(r["solidworks_simulation"] is not None for r in rows)

    # ---- title and summary on the first page
    _para(doc, "AsmFEA verification report", 24, True, BLUE)
    _para(doc, "Classical-mechanics checks of the full chain, SolidWorks to report, and of the CalculiX solver", 12,
          colour=GREY)
    _kv_table(doc, [["Date", f"{when:%d %B %Y}"], ["AsmFEA", versions["asmfea"]],
                    ["Solver", f"CalculiX {versions['calculix']} (MD5 {versions.get('calculix_md5', '-')})"],
                    ["Source of results", "Full-chain runs from SolidWorks models" if result["source"] == "sw"
                     else "Regression runs from the stored neutral models (no SolidWorks)"],
                    ["Overall", "All comparisons within their acceptance criteria" if result["all_pass"]
                     else "One or more comparisons outside their acceptance criteria"]], (4.5, 12.5), size=9.5)
    doc.add_heading("1. Summary", 1)
    head = ["Case", "Quantity", "Reference", "AsmFEA", "Difference", "Result"]
    if sws_any:
        head.insert(5, "SOLIDWORKS Simulation")
    trows = []
    for r in rows:
        row = [r["case"], f"{r['quantity']} ({r['unit']})", _num(r["reference"]), _num(r["asmfea"]),
               _pct(r["difference"]), PASS_FAIL[r["pass"]]]
        if sws_any:
            sws = r["solidworks_simulation"]
            row.insert(5, f"{_num(sws)} ({_pct(r['sws_difference'])})" if sws is not None else
                       ("not available" if r["criterion"] == "buckling" else "not supplied"))
        trows.append(row)
    _table(doc, head, trows, [1.2, 5.2, 2.3, 2.3, 2.0, 1.4] if not sws_any else [1.1, 4.4, 2.0, 2.0, 1.8, 2.8, 1.3],
           size=8)
    npass = sum(r["pass"] for r in rows)
    _para(doc, f"{npass} of {len(rows)} comparisons pass. The reference is the classical solution, except "
               "for the mass rows marked SolidWorks, which use SolidWorks' mass properties. Difference = (AsmFEA - "
               "reference) / reference, sign kept. Criteria: deflection, stress and buckling load within 5 %, "
               "reactions within 1 %, mass within 0.5 %.", 9)
    if not sws_any:
        _para(doc, "SOLIDWORKS Simulation results for the static cases are to be added by the user (SOLIDWORKS "
                   "Premium, linear static). " + NO_SWS_BUCKLING + ": for C4 the comparison is with theory alone; "
                   "AsmFEA can run this analysis and SOLIDWORKS Premium cannot.", 9, colour=GREY)
    worst = max(rows, key=lambda r: abs(r["difference"]))
    _para(doc, "Conclusion. On these cases AsmFEA reproduces the classical answers for deflection, stress, "
               f"reactions, mass and buckling load; the largest difference is {_pct(worst['difference'])} "
               f"({worst['case']}, {worst['quantity'].lower()}). The solver itself reproduces the reference results of "
               "its own test examples and the NAFEMS LE1 and LE10 targets (section 3). These results show that the "
               "chain poses and reports these models correctly; they are not a certification (section 7).", 10)

    _para(doc, USER_NOTE, 9.5, True, colour=BLUE)

    # ---- 2. How the analysis works
    doc.add_heading("2. How the analysis works", 1)
    _para(doc, T.NOT_STRESS, 10, True)
    doc.add_picture(str(_diagram(work / "diagram.png")), width=Cm(16.5))
    _para(doc, "Two things are therefore shown separately: that the solver solves a correctly posed model accurately "
               "(section 3, CalculiX alone), and that AsmFEA poses each model correctly and reports it correctly "
               "(sections 5 and 6, the full chain from SolidWorks).", 9.5)

    # ---- 3. The solver
    doc.add_heading("3. The solver", 1)
    sv = result.get("solver") or {}
    _para(doc, "CalculiX is a free, open-source finite element program by Guido Dhondt (solver, ccx) and Klaus "
               "Wittig (pre- and post-processor), distributed under the GNU GPL. AsmFEA runs it as a separate "
               "program. CalculiX has no formal certification and no vendor verification manual of the kind "
               "published for the large commercial codes; it is not described here as certified or validated by "
               "any body. What exists is its developer's own test examples and published benchmark comparisons, "
               "and both were run on the exact binary AsmFEA uses.", 9.5)
    if sv:
        ex, pa = sv.get("examples", {}), sv.get("examples", {}).get("pardiso_pass", {})
        rows = [["Binary", f"{sv.get('binary', '-')} ({sv.get('ccx_version', 'CalculiX 2.22')})"],
                ["MD5", f"{sv.get('md5', '-')} (the same file as in the author's calculix_2.22_4win.zip, which the "
                        "AsmFEA installer downloads)"],
                ["Official test examples", f"{ex.get('matched')} of {ex.get('run')} decks match their reference results "
                                           "(official compare scripts; a value fails beyond 0.1 % of the largest value "
                                           "in its output block)"],
                ["With PARDISO (AsmFEA's solver)", f"{pa.get('matched', '-')} of {pa.get('run', '-')}; the exception "
                                                   "tests the iterative Cholesky solver, whose reference values carry "
                                                   "that solver's convergence error"]]
        for key, name, cols in (("le1", "NAFEMS LE1", "own_decks_cps8"), ("le10", "NAFEMS LE10", "own_decks_c3d10")):
            le = sv.get(key, {})
            decks = le.get(cols) or []
            if decks:
                f, t = decks[-1], le["target_MPa"]
                size = f"{f['elements']:,} {f['element'].split()[0]} elements" + (
                    f", {f['equations']:,} equations" if f.get("equations") else "")
                rows.append([f"{name} (target {t:g} MPa)",
                             f"finest mesh {f['syy_D_MPa']:.4g} MPa, {_pct((f['syy_D_MPa'] - t) / t)} ({size}); "
                             f"{len(decks)} meshes converging on the target"])
        _kv_table(doc, rows, (5.0, 12.0), size=8.5)
        _para(doc, "Full detail, decks and scripts: cases/verification/solver_checks (SUMMARY.md, results.json). The "
                   "official set is thin on quadratic tetrahedra, wedges, gravity and buckling; the LE10 runs and the "
                   "full-chain cases below cover those.", 8.5, colour=GREY)
    _table(doc, ["Source", "What it shows", "Link"], [[a, b, c or "-"] for a, b, c in SOURCES], [4.5, 6.0, 6.5],
           size=7.5)
    _para(doc, "All links were checked on the date of this report. The bConverged pages compare CalculiX 1.5 (2006); "
               "the NAFEMS targets are quoted from NAFEMS TNSB Rev. 3 (1990) via bConverged, as the NAFEMS document "
               "itself is not freely available.", 8, colour=GREY)

    # ---- 3b. The hybrid solver for large models
    doc.add_heading("3b. " + T.HYBRID_TITLE, 1)
    _para(doc, T.HYBRID_INTRO, 9.5)
    for line in T.HYBRID_HOW:
        doc.add_paragraph(line, style="List Bullet")
    _table(doc, ["Model", "Result"], [[a, b] for a, b in T.HYBRID_EVIDENCE], [5.5, 11.5], size=8)
    _para(doc, T.HYBRID_NOTE, 8.5, colour=GREY)

    # ---- 4. Method
    doc.add_heading("4. Method", 1)
    for t in ["Full chain: every case starts as a SolidWorks model and runs through the add-in, the pipeline and the "
              "report generator, with nothing edited by hand in between. Loads and fixtures were set through the "
              "add-in on faces selected in SolidWorks.",
              "Theory first: each classical calculation was written and committed before the analysis for that case "
              "was run, so neither could be adjusted to agree with the other (git history of cases/verification).",
              "Like with like: slender members (length at least 20 times depth); section properties from the "
              "modelled section with its corner radii; shear deflection included; stresses read at least one section "
              "depth clear of supports and load points; nominal stress compared with nominal stress.",
              "Default settings: every case at the default mesh fidelity (standard), which runs a base and a refined "
              "mesh and reports the refined one; C1 also at coarse and fine (section 6).",
              "Read-outs: each value is the mean of CalculiX's nodal results over a node set defined by coordinates, "
              "stated with each result; global values (mass, reaction, buckling factor) come from AsmFEA's results.",
              "No tuning: no mesh setting was changed per case to obtain a pass.",
              "SOLIDWORKS Simulation, where supplied, is a second opinion, not an acceptance criterion: two finite "
              "element codes agreeing is weaker evidence than agreement with theory, since they share the same kind "
              "of assumptions."]:
        _para(doc, "• " + t, 9.5)
    _table(doc, ["Quantity", "Criterion"], [["Deflection and rotation", "within 5 % of theory"],
                                           ["Nominal stress at the stated read-out point", "within 5 %"],
                                           ["Buckling load", "within 5 %"],
                                           ["Reactions", "within 1 % of applied load"],
                                           ["Mass", "within 0.5 % of the SolidWorks mass properties"]], [8.0, 9.0])

    # ---- 5. Cases
    doc.add_heading("5. Cases", 1)
    for cr in result["cases"]:
        c, m = cr["case"], cr["measured"]
        cdir = Path(c["dir"])
        doc.add_heading(f"{c['id']}  {c['title']}", 2)
        _kv_table(doc, [["Exercises", c["exercises"]], ["Model", c["description"]], ["Restraints and loads", c["restraints"]],
                        ["Mesh", f"default fidelity (standard); reported mesh: {m['level']}, {m['dof']:,} degrees of "
                                 f"freedom"],
                        ["Convergence check", f"largest change between base and refined meshes "
                                              f"{_pct(m['convergence_worst'])}" +
                                              (" (flagged CONVERGENCE_POOR: see the comment)" if "CONVERGENCE_POOR" in
                                               m["warnings"] else "")]], (4.0, 13.0), size=8.5)
        res = Path(m["study"]) / "results"
        if (res / "setup_iso.png").exists():
            _picture(doc, res / "setup_iso.png", "Model as analysed: fixtures blue, loaded faces red", 8.0)
        _para(doc, "Classical calculation (written before the run):", 9.5, True)
        _mono(doc, _theory_text(cdir))
        _para(doc, "Comparison:", 9.5, True)
        _table(doc, ["Quantity", "Reference", "AsmFEA", "Difference", "Result", "Read-out"],
               [[f"{r['quantity']} ({r['unit']})", _num(r["reference"]), _num(r["asmfea"]), _pct(r["difference"]),
                 PASS_FAIL[r["pass"]], r["where"] + (f" ({r['nodes']} nodes)" if r["nodes"] else "")]
                for r in cr["rows"]], [3.6, 2.0, 2.0, 1.8, 1.2, 6.4], size=7.5)
        plot = next((res / i for i in m["images"] if "buckling_mode1_iso" in i), None) or \
            next((res / i for i in m["images"] if i.endswith("_vonmises_iso.png")), None)
        if plot is not None and plot.exists():
            _picture(doc, plot, "Result: " + ("buckling mode 1" if "buckling" in plot.name else "von Mises stress"), 11.0)
        notes = T.CASE_NOTES
        _para(doc, "Comment. " + notes.get(c["id"], ""), 9.5)

    # ---- 5b. Methods for joints and notches
    doc.add_heading("5b. Methods for joints and notches", 1)
    _para(doc, T.METHODS_INTRO, 9.5)
    _table(doc, ["Method", "What it does", "Situations it is there for", "Evidence"],
           [list(m) for m in T.JOINT_METHODS], [3, 6, 4.5, 4.5])
    if result.get("methods"):
        _para(doc, "Method checks: hand-written models run through the pipeline (no SolidWorks), each against a hand "
                   "answer. M4 checks the receiving-wall method, which is not used in normal runs (see Limitations).", 9.5)
        for mc in result["methods"]:
            _para(doc, f"{mc['id']}  {mc['title']}", 10, True)
            _table(doc, ["Quantity", "Reference", "AsmFEA", "Difference", "Result"],
                   [[r["quantity"], _num(r["reference"]), _num(r["asmfea"]), _pct(r["difference"]),
                     "PASS" if r["pass"] else "FAIL"] for r in mc["rows"]], [8, 2.5, 2.5, 2.5, 1.8])

    # ---- 6. Mesh sensitivity
    doc.add_heading("6. Mesh sensitivity", 1)
    ms = [m for m in result["mesh_sensitivity"] if m["case"] == "C1"]
    if ms:
        order = {"coarse": 0, "standard": 1, "fine": 2}
        ms.sort(key=lambda m: order.get(m["fidelity"], 9))
        t = []
        for m in ms:
            r1 = next(r for r in m["rows"] if r["quantity"] == "Tip deflection")
            r2 = next(r for r in m["rows"] if r["quantity"].startswith("Bending stress"))
            t.append([m["fidelity"], f"{m['measured']['dof']:,}", _num(r1["asmfea"]), _pct(r1["difference"]),
                      _num(r2["asmfea"]), _pct(r2["difference"]), _pct(m["measured"]["convergence_worst"]),
                      "yes" if "CONVERGENCE_POOR" in m["measured"]["warnings"] else "no"])
        _table(doc, ["Fidelity", "DOF", "Tip (mm)", "vs theory", "σ mid (MPa)", "vs theory", "Check: largest change",
                     "Flagged"], t, [1.8, 2.0, 1.8, 1.8, 2.0, 1.8, 2.8, 1.5], size=8)
        if len(ms) >= 2:
            doc.add_picture(str(_mesh_plot(ms, work / "mesh_c1.png")), width=Cm(16.5))
        _para(doc, T.MESH_NOTE, 9.5)
    else:
        _para(doc, "Not run yet.", 9.5)

    # ---- 7. Limitations
    doc.add_heading("7. Limitations", 1)
    for t in ["Linear elastic material, small displacements, linear buckling (no imperfections, no plasticity).",
              "Not covered by a case: separating contact and friction, preloaded bolts, revolute and sliding joints, "
              "weld beads, sheet-metal shells meshed through the thickness, pressure loads, remote torque, stress "
              "concentrations (the extended cases E1-E5 of the plan).",
              "Fatigue and dynamics are not part of AsmFEA's first release and are not verified.",
              "Four cases cannot cover every geometry; they test the parts of the chain listed under each case.",
              "Neither CalculiX nor AsmFEA is certified or validated by any body. This report compares two methods "
              "on a few cases; it is not an approval of any particular analysis. " + USER_NOTE]:
        _para(doc, "• " + t, 9.5)

    # ---- 8. Appendices
    doc.add_heading("8. Appendices", 1)
    doc.add_heading("A. Software versions", 2)
    first = result["cases"][0]["measured"] if result["cases"] else {}
    sw_version = "-"
    try:
        sw_version = json.loads((Path(first["study"]) / "model.json").read_text(encoding="utf-8-sig"))["source"].get(
            "sw_version", "-")
        major = int(str(sw_version).split(".")[0])
        sw_version = f"SOLIDWORKS {major + 1992} ({sw_version})"
    except Exception:
        pass
    _kv_table(doc, [["AsmFEA", versions["asmfea"]], ["CalculiX", f"{versions['calculix']} (ccx_dynamic.exe, PARDISO via "
                                                                  f"Intel MKL 2022.2.1), MD5 {versions.get('calculix_md5', '-')}"],
                    ["Gmsh", versions.get("gmsh", "-")], ["SolidWorks", f"{sw_version}"],
                    ["Python", platform.python_version()], ["Machine", platform.platform()]], (4.0, 13.0), size=8.5)
    doc.add_heading("B. Inputs and generated CalculiX files", 2)
    hrows = []
    for cr in result["cases"]:
        st = Path(cr["measured"]["study"])
        for f in [st / "model.json", st / "study.json", *sorted((st / "parts").glob("*.step")), Path(cr["measured"]["inp"])]:
            if f.exists():
                hrows.append([cr["case"]["id"], str(f.relative_to(root)) if root in f.parents else str(f),
                              f"{f.stat().st_size:,}", _sha256(f)[:16]])
    _table(doc, ["Case", "File", "Bytes", "SHA-256 (first 16)"], hrows, [1.2, 9.6, 2.0, 4.2], size=7)
    _para(doc, "The CalculiX input files (study.inp) are kept in each study's run folder; they are too large to print.",
          8, colour=GREY)
    doc.add_heading("C. Findings during verification", 2)
    for t in ["C2's first run was refused by AsmFEA's mechanism check: two plain pinned ends leave a bar free to spin "
              "about its own axis. The model was corrected (twist held at one end); the classical calculation did not "
              "change.",
              "C2 showed that AsmFEA reported the reaction from CalculiX's RF output, which leaves out load applied "
              "directly to held nodes (here 0.10 %, the self-weight on the supported faces). AsmFEA now also reports "
              "support_reaction, the full load the supports carry, and the report's resultant forces use it. The "
              "analysis itself was not affected: its balance check already included that load.",
              "The automation calls used to set up and capture these models arrived on a background thread; the "
              "add-in now hands pane work to its own thread."]:
        _para(doc, "• " + t, 9)
    doc.add_heading("D. Calculation scripts", 2)
    _table(doc, ["Case", "Script", "SHA-256 (first 16)"],
           [[cr["case"]["id"], str((Path(cr["case"]["dir"]) / "theory.py").relative_to(root)),
             _sha256(Path(cr["case"]["dir"]) / "theory.py")[:16]] for cr in result["cases"]], [1.5, 9.0, 6.5], size=8)
    doc.save(str(out))
    return out
